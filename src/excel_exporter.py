# src/excel_exporter.py
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment


# 项目根目录（src 的上一级）
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_OUTPUT_DIR = os.path.join(PROJECT_ROOT, "output")


class ExcelExporter:
    @staticmethod
    def export(papers, filename=None, output_dir=DEFAULT_OUTPUT_DIR):
        """
        将论文列表写入 Excel。
        :param papers: 论文列表
        :param filename: 输出文件名（None 则自动生成带时间戳的名字）
        :param output_dir: 输出目录
        :return: 成功时返回完整文件路径，失败返回 None
        """
        if not papers:
            print("[Excel] 没有数据可导出。")
            return None

        os.makedirs(output_dir, exist_ok=True)

        if filename is None:
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"arxiv_results_{ts}.xlsx"

        filepath = os.path.join(output_dir, filename)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ArXiv Papers"

        # 表头
        headers = ["标题", "摘要", "arXiv 链接", "发布日期"]
        ws.append(headers)

        header_font = Font(bold=True)
        for col_num, _ in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        # 数据
        for paper in papers:
            ws.append([
                paper['title'],
                paper['summary'],
                paper['link'],
                paper['published'],
            ])

        # 列宽
        ws.column_dimensions['A'].width = 50
        ws.column_dimensions['B'].width = 80
        ws.column_dimensions['C'].width = 35
        ws.column_dimensions['D'].width = 15

        # 摘要列自动换行
        for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')

        # arXiv 链接做成超链接
        for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
            for cell in row:
                if cell.value:
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0563C1", underline="single")

        try:
            wb.save(filepath)
            print(f"[Excel] 已写入: {filepath}")
            return filepath
        except Exception as e:
            print(f"[Excel] 写入失败: {e}")
            return None